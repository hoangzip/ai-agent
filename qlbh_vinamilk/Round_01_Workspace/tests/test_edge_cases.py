"""
Module: tests/test_edge_cases.py
Kiểm thử tự động các trường hợp biên (Edge Cases), dữ liệu lỗi, phân loại đặc biệt
và chính sách Deduplication/Conflict theo yêu cầu bắt buộc của Round 01.
"""

import os
import pytest
import openpyxl
from src.parser import InventoryExcelParser, ExcelParseError, ExcelValidationError, EXPECTED_HEADERS
from src.comparator import InventoryComparator
from src.identity import SnapshotIdentity

FIXTURE_DIR = "fixtures"
os.makedirs(FIXTURE_DIR, exist_ok=True)


def create_synthetic_excel(filepath: str, headers=None, summary_qty=10.0, summary_val=100000.0, rows=None):
    """Tạo file Excel synthetic cho mục đích test."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kiem-ke"

    # Header
    hdrs = headers if headers is not None else EXPECTED_HEADERS
    for c_idx, h in enumerate(hdrs, start=1):
        ws.cell(1, c_idx, value=h)

    # Summary row (row 2)
    ws.cell(2, 5, value=summary_qty)
    ws.cell(2, 8, value=summary_val)
    ws.cell(2, 11, value="=SUM(K3:K10)")
    ws.cell(2, 12, value="=K2-E2")

    # Detail rows (row 3 onwards)
    if rows:
        for r_idx, r_data in enumerate(rows, start=3):
            for c_idx, val in enumerate(r_data, start=1):
                ws.cell(r_idx, c_idx, value=val)

    wb.save(filepath)
    return filepath


def test_edge_division_by_zero_and_zero_quantity():
    """
    Edge case: Quantity A = 0.
    Phần trăm biến động phải không gây lỗi ZeroDivisionError (trả về None / an toàn).
    """
    snap_a = {
        "metadata": {"filename": "snap_a.xlsx"},
        "items": [
            {
                "product_code": "01SKU01",
                "warehouse": "Kho hàng bán",
                "product_name": "Sản phẩm A",
                "quantity": 0.0,
                "unit_price": 50000.0,
                "total_amount": 0.0
            }
        ]
    }
    snap_b = {
        "metadata": {"filename": "snap_b.xlsx"},
        "items": [
            {
                "product_code": "01SKU01",
                "warehouse": "Kho hàng bán",
                "product_name": "Sản phẩm A",
                "quantity": 10.0,
                "unit_price": 50000.0,
                "total_amount": 500000.0
            }
        ]
    }

    cmp = InventoryComparator(snap_a, snap_b).compare()
    assert cmp["summary"]["matched_count"] == 1
    assert cmp["summary"]["increased_count"] == 1
    detail = cmp["details"][0]
    assert detail["diff_qty"] == 10.0
    # Phép chia cho 0 được xử lý an toàn
    assert detail["pct_qty_change"] is None
    assert detail["pct_val_change"] is None


def test_edge_missing_and_new_sku_classification():
    """
    Edge case: SKU có trong A nhưng không có trong B -> NOT_APPEARED (không tự gán B=0).
    SKU có trong B nhưng không có trong A -> NEW.
    """
    snap_a = {
        "metadata": {"filename": "snap_a.xlsx"},
        "items": [
            {
                "product_code": "SKU_EXIST_A",
                "warehouse": "Kho hàng bán",
                "product_name": "SP A",
                "quantity": 5.0,
                "unit_price": 10000.0,
                "total_amount": 50000.0
            }
        ]
    }
    snap_b = {
        "metadata": {"filename": "snap_b.xlsx"},
        "items": [
            {
                "product_code": "SKU_NEW_B",
                "warehouse": "Kho hàng bán",
                "product_name": "SP B",
                "quantity": 8.0,
                "unit_price": 20000.0,
                "total_amount": 160000.0
            }
        ]
    }

    cmp = InventoryComparator(snap_a, snap_b).compare()
    summary = cmp["summary"]
    assert summary["matched_count"] == 0
    assert summary["new_count"] == 1
    assert summary["not_appeared_count"] == 1

    details = {d["product_code"]: d for d in cmp["details"]}
    assert details["SKU_EXIST_A"]["status"] == "NOT_APPEARED"
    assert details["SKU_EXIST_A"]["qty_b"] is None  # Không tự suy luận tồn B = 0

    assert details["SKU_NEW_B"]["status"] == "NEW"
    assert details["SKU_NEW_B"]["qty_a"] is None


def test_edge_price_changed_quantity_unchanged():
    """
    Edge case: Số lượng tồn giữ nguyên nhưng đơn giá sau VAT thay đổi.
    Thành tiền thay đổi, diff_qty = 0, phát hiện trong price_changed_qty_same.
    """
    snap_a = {
        "metadata": {"filename": "snap_a.xlsx"},
        "items": [
            {
                "product_code": "SKU_PRICE_CHG",
                "warehouse": "Kho hàng bán",
                "product_name": "SP Đổi Giá",
                "quantity": 10.0,
                "unit_price": 20000.0,
                "total_amount": 200000.0
            }
        ]
    }
    snap_b = {
        "metadata": {"filename": "snap_b.xlsx"},
        "items": [
            {
                "product_code": "SKU_PRICE_CHG",
                "warehouse": "Kho hàng bán",
                "product_name": "SP Đổi Giá",
                "quantity": 10.0,
                "unit_price": 25000.0,
                "total_amount": 250000.0
            }
        ]
    }

    cmp = InventoryComparator(snap_a, snap_b).compare()
    assert cmp["summary"]["matched_count"] == 1
    assert cmp["summary"]["unchanged_count"] == 1
    assert cmp["summary"]["total_delta_quantity"] == 0.0
    assert cmp["summary"]["total_delta_value"] == 50000.0
    assert cmp["summary"]["price_changed_qty_same_count"] == 1


def test_edge_same_sku_in_two_different_warehouses():
    """
    Edge case: Cùng 1 mã SKU xuất hiện ở 2 kho khác nhau -> được coi là 2 bản ghi độc lập.
    """
    snap_a = {
        "metadata": {"filename": "snap_a.xlsx"},
        "items": [
            {
                "product_code": "02MULTI",
                "warehouse": "Kho hàng bán",
                "product_name": "SP Multi",
                "quantity": 10.0,
                "unit_price": 10000.0,
                "total_amount": 100000.0
            },
            {
                "product_code": "02MULTI",
                "warehouse": "Kho hàng vật phẩm KM",
                "product_name": "SP Multi",
                "quantity": 3.0,
                "unit_price": 0.0,
                "total_amount": 0.0
            }
        ]
    }
    snap_b = {
        "metadata": {"filename": "snap_b.xlsx"},
        "items": [
            {
                "product_code": "02MULTI",
                "warehouse": "Kho hàng bán",
                "product_name": "SP Multi",
                "quantity": 15.0,
                "unit_price": 10000.0,
                "total_amount": 150000.0
            },
            {
                "product_code": "02MULTI",
                "warehouse": "Kho hàng vật phẩm KM",
                "product_name": "SP Multi",
                "quantity": 2.0,
                "unit_price": 0.0,
                "total_amount": 0.0
            }
        ]
    }

    cmp = InventoryComparator(snap_a, snap_b).compare()
    assert cmp["summary"]["matched_count"] == 2
    assert cmp["summary"]["increased_count"] == 1  # Kho hàng bán tăng từ 10 -> 15
    assert cmp["summary"]["decreased_count"] == 1  # Kho hàng KM giảm từ 3 -> 2


def test_validation_error_duplicate_sku_in_same_warehouse():
    """
    Validation error: Trùng mã SKU trong CÙNG một kho trong 1 file.
    Hệ thống bắt buộc phải báo lỗi validation, không được tự ý ghi đè.
    """
    file_dup = os.path.join(FIXTURE_DIR, "synthetic_dup_sku.xlsx")
    rows = [
        [1, "A", "02DUP", "SP Trùng", 5.0, "Kho hàng bán", 10000.0, 50000.0, 0, 0, "=0", "=0", "THU", 1, "Hop"],
        [2, "A", "02DUP", "SP Trùng", 5.0, "Kho hàng bán", 10000.0, 50000.0, 0, 0, "=0", "=0", "THU", 1, "Hop"]
    ]
    create_synthetic_excel(file_dup, summary_qty=10.0, summary_val=100000.0, rows=rows)

    parser = InventoryExcelParser(file_dup)
    with pytest.raises(ExcelValidationError, match="Trùng lặp khóa tự nhiên"):
        parser.parse()


def test_validation_error_negative_quantity():
    """
    Validation error: SL tồn kho âm.
    """
    file_neg = os.path.join(FIXTURE_DIR, "synthetic_neg_qty.xlsx")
    rows = [
        [1, "A", "02NEG", "SP Âm", -5.0, "Kho hàng bán", 10000.0, -50000.0, 0, 0, "=0", "=0", "THU", 1, "Hop"]
    ]
    create_synthetic_excel(file_neg, summary_qty=-5.0, summary_val=-50000.0, rows=rows)

    parser = InventoryExcelParser(file_neg)
    with pytest.raises(ExcelValidationError, match="SL tồn kho không được âm"):
        parser.parse()


def test_validation_error_missing_header():
    """
    Parse error: File thiếu hoặc sai cấu trúc header 15 cột.
    """
    file_bad_hdr = os.path.join(FIXTURE_DIR, "synthetic_bad_header.xlsx")
    bad_headers = EXPECTED_HEADERS[:10]  # chỉ có 10 cột
    create_synthetic_excel(file_bad_hdr, headers=bad_headers)

    parser = InventoryExcelParser(file_bad_hdr)
    with pytest.raises(ExcelParseError, match="Cấu trúc 15 cột không khớp"):
        parser.parse()


def test_validation_error_summary_mismatch():
    """
    Validation error: Tổng dòng chi tiết lệch so với số khai báo ở dòng 2.
    """
    file_mismatch = os.path.join(FIXTURE_DIR, "synthetic_summary_mismatch.xlsx")
    rows = [
        [1, "A", "02MIS", "SP Lệch", 10.0, "Kho hàng bán", 10000.0, 100000.0, 0, 0, "=0", "=0", "THU", 1, "Hop"]
    ]
    # Dòng 2 khai báo 15.0 nhưng chi tiết chỉ có 10.0
    create_synthetic_excel(file_mismatch, summary_qty=15.0, summary_val=100000.0, rows=rows)

    parser = InventoryExcelParser(file_mismatch)
    with pytest.raises(ExcelValidationError, match="Lệch tổng số lượng tồn"):
        parser.parse()


def test_identity_dedup_and_conflict_policies():
    """
    Kiểm tra toàn bộ ma trận chính sách tiếp nhận snapshot:
    NEW -> DUPLICATE -> CONFLICT -> ALREADY_IMPORTED -> SOFT_DELETED
    """
    mock_items_v1 = [
        {"product_code": "02TEST", "warehouse": "Kho hàng bán", "quantity": 10.0, "unit_price": 5000.0, "total_amount": 50000.0}
    ]
    mock_items_v2 = [
        {"product_code": "02TEST", "warehouse": "Kho hàng bán", "quantity": 12.0, "unit_price": 5000.0, "total_amount": 60000.0}
    ]

    id1 = SnapshotIdentity("KiemKe_KK.CH41551.26.339000_20261009.xlsx", file_sha256="binhash1", parsed_items=mock_items_v1)
    meta1 = id1.to_metadata_dict()

    db_records = {}

    # 1. NEW
    action, reason = SnapshotIdentity.evaluate_ingestion_policy(meta1, db_records)
    assert action == "NEW"

    # Giả lập đã lưu vào DB
    db_records[meta1["full_code"]] = {
        "full_code": meta1["full_code"],
        "file_sha256": "binhash1",
        "content_hash": meta1["content_hash"],
        "is_deleted": False
    }

    # 2. ALREADY_IMPORTED (cùng file_sha256)
    action, _ = SnapshotIdentity.evaluate_ingestion_policy(meta1, db_records)
    assert action == "ALREADY_IMPORTED"

    # 3. DUPLICATE (khác file binary sha256 do re-save, nhưng nội dung y hệt)
    meta_resaved = meta1.copy()
    meta_resaved["file_sha256"] = "binhash_different_resaved"
    action, _ = SnapshotIdentity.evaluate_ingestion_policy(meta_resaved, db_records)
    assert action == "DUPLICATE"

    # 4. CONFLICT (cùng mã kiểm kê nhưng nội dung thay đổi)
    id_conflict = SnapshotIdentity("KiemKe_KK.CH41551.26.339000_20261009.xlsx", file_sha256="binhash_modified", parsed_items=mock_items_v2)
    meta_conflict = id_conflict.to_metadata_dict()
    action, reason = SnapshotIdentity.evaluate_ingestion_policy(meta_conflict, db_records)
    assert action == "CONFLICT"
    assert "Cấm tự động ghi đè" in reason

    # 5. SOFT_DELETED (chứng từ đã bị xóa mềm, không tự động phục hồi)
    db_records[meta1["full_code"]]["is_deleted"] = True
    action, reason = SnapshotIdentity.evaluate_ingestion_policy(meta1, db_records)
    assert action == "SOFT_DELETED"
    assert "Không tự động phục hồi" in reason
